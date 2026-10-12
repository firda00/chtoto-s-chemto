from __future__ import annotations

from pathlib import Path

import pytest

from ats_tailor.parsers import (
    detect_heading,
    detect_language,
    extract_contacts,
    extract_text,
    load_document,
    load_resume,
    parse_resume,
    split_sections,
)
from conftest import make_simple_pdf

EN_LINES = [
    "John Doe",
    "john.doe@example.com | +1 555 010 0100",
    "SUMMARY",
    "Backend engineer with 3 years of experience in Python.",
    "SKILLS",
    "Python, FastAPI, Docker",
    "EXPERIENCE",
    "Acme Corp - Python Developer",
    "Jan 2021 - Present",
    "EDUCATION",
    "BSc Computer Science, 2016 - 2020",
]


def test_txt_and_md_extraction(tmp_path: Path) -> None:
    txt = tmp_path / "cv.txt"
    txt.write_text("Навыки\nPython, SQL\n", encoding="utf-8")
    assert "Python" in extract_text(txt)

    md = tmp_path / "cv.md"
    md.write_text(
        "# Имя\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n![photo](me.png)\n", encoding="utf-8"
    )
    doc = load_document(md)
    assert doc.fmt == "md"
    assert doc.tables == 1
    assert doc.images == 1


def test_cp1251_txt_is_decoded(tmp_path: Path) -> None:
    p = tmp_path / "old.txt"
    p.write_bytes("Опыт работы\nООО Ромашка".encode("cp1251"))
    assert "Ромашка" in extract_text(p)


def test_docx_extraction_with_table_and_bullets(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Иван Петров")
    doc.add_paragraph("ivan@example.com, +7 900 111-22-33")
    doc.add_heading("Опыт работы", level=1)
    doc.add_paragraph("Разработал RAG-сервис на FastAPI", style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Docker"
    path = tmp_path / "cv.docx"
    doc.save(str(path))

    info = load_document(path)
    assert info.fmt == "docx"
    assert info.tables == 1
    assert "- Разработал RAG-сервис на FastAPI" in info.text
    assert "Python | Docker" in info.text

    resume = load_resume(path)
    assert "RAG-сервис" in resume.section("experience")
    assert resume.contacts.emails == ["ivan@example.com"]


def test_pdf_extraction(tmp_path: Path) -> None:
    try:
        path = make_simple_pdf(tmp_path / "cv.pdf", EN_LINES)
    except Exception as exc:  # pragma: no cover - зависит от версии pypdf
        pytest.skip(f"не удалось сгенерировать PDF: {exc}")
    info = load_document(path)
    assert info.fmt == "pdf"
    assert info.pages == 1
    assert "Python" in info.text
    resume = load_resume(path)
    assert "FastAPI" in resume.section("skills")
    assert resume.contacts.emails == ["john.doe@example.com"]


def test_unsupported_format(tmp_path: Path) -> None:
    p = tmp_path / "cv.exe"
    p.write_bytes(b"MZ")
    with pytest.raises(ValueError):
        load_document(p)


def test_split_sections_ru_sample(resume_ru) -> None:
    for key in ("contacts", "summary", "skills", "experience", "education"):
        assert resume_ru.section(key), key
    assert "Нейросфера" in resume_ru.section("experience")
    assert "pgvector" in resume_ru.section("skills")
    assert "languages" in resume_ru.sections
    assert resume_ru.contacts.name == "Алексей Смирнов"
    assert resume_ru.contacts.emails == ["alexey.smirnov@example.com"]
    assert resume_ru.contacts.phones
    assert resume_ru.language == "ru"


def test_split_sections_en_sample(resume_en) -> None:
    for key in ("contacts", "summary", "skills", "experience", "education"):
        assert resume_en.section(key), key
    assert "Brightline" in resume_en.section("experience")
    assert resume_en.contacts.name == "Maria Ivanova"
    assert resume_en.language == "en"


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("## Опыт работы", "experience"),
        ("ОПЫТ РАБОТЫ", "experience"),
        ("Опыт работы — 6 лет 2 месяца", "experience"),
        ("**Ключевые навыки:**", "skills"),
        ("Work Experience", "experience"),
        ("TECHNICAL SKILLS", "skills"),
        ("Professional Summary", "summary"),
        ("О себе", "summary"),
        ("Education", "education"),
        ("Контакты", "contacts"),
        ("Опыт работы с LLM более 3 лет.", None),
        ("Разработал RAG-систему поиска", None),
    ],
)
def test_detect_heading(line: str, expected: str | None) -> None:
    assert detect_heading(line) == expected


def test_inline_heading_content() -> None:
    sections = split_sections("Иван\nНавыки: Python, SQL\nОбразование: МГУ, 2015")
    assert sections["skills"] == "Python, SQL"
    assert sections["education"] == "МГУ, 2015"
    assert sections["contacts"] == "Иван"


def test_no_headings_goes_to_other() -> None:
    assert split_sections("просто текст без заголовков") == {"other": "просто текст без заголовков"}


def test_contacts_extraction() -> None:
    c = extract_contacts(
        "a.b@example.com, +998 90 123 45 67, github.com/demo, Telegram: @demo_user\n2019 - 2021"
    )
    assert c.emails == ["a.b@example.com"]
    assert c.phones == ["+998 90 123 45 67"]
    assert "github.com/demo" in c.links
    assert "@demo_user" in c.links
    # диапазон лет не телефон, домен email не ссылка
    assert all("2019" not in p for p in c.phones)
    assert "@example" not in c.links


def test_contacts_at_bottom_without_heading() -> None:
    text = "Навыки\nPython\n\nОпыт работы\nООО Ромашка 2020 — 2022\nmail@example.com"
    r = parse_resume(text)
    assert r.contacts.emails == ["mail@example.com"]
    assert "contacts" in r.sections


def test_detect_language() -> None:
    assert detect_language("Опыт работы с Python и Docker") == "ru"
    assert detect_language("Experience with Python and Docker") == "en"
