from __future__ import annotations

from pathlib import Path

from ats_tailor.export import (
    changelog_markdown,
    to_markdown,
    write_all,
    write_cover_letter,
    write_docx,
    write_markdown,
)
from ats_tailor.llm import MockLLM
from ats_tailor.parsers import load_resume
from ats_tailor.tailor import cover_letter, tailor_resume


def test_markdown_and_docx_export(tmp_path: Path, resume_ru, vacancy_ru) -> None:
    tr = tailor_resume(resume_ru, vacancy_ru, MockLLM(), platform="hh")

    md = to_markdown(tr)
    assert md.startswith("# Алексей Смирнов")
    assert "## Ключевые навыки" in md
    assert "## Опыт работы" in md

    md_path = write_markdown(tr, tmp_path / "out" / "cv.md")
    assert md_path.read_text(encoding="utf-8") == md

    docx_path = write_docx(tr, tmp_path / "out" / "cv.docx")
    from docx import Document

    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    assert "Алексей Смирнов" in texts
    assert "Опыт работы" in texts
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert len(doc.tables) == 0  # ATS-friendly: без таблиц

    # экспорт снова читается нашим парсером и делится на разделы
    back = load_resume(docx_path)
    assert "Нейросфера" in back.section("experience")
    assert back.section("skills")


def test_linkedin_titles_en(tmp_path: Path, resume_en, vacancy_en) -> None:
    tr = tailor_resume(resume_en, vacancy_en, MockLLM(), platform="linkedin")
    md = to_markdown(tr)
    assert md.startswith("## About")
    assert "## Contact" in md


def test_write_all_and_changelog(tmp_path: Path, resume_ru, vacancy_ru) -> None:
    tr = tailor_resume(resume_ru, vacancy_ru, MockLLM(), platform="generic")
    paths = write_all(tr, tmp_path, stem="cv")
    assert [p.name for p in paths] == [
        "cv_generic.md",
        "cv_generic.docx",
        "cv_generic_changelog.md",
    ]
    assert all(p.exists() and p.stat().st_size > 0 for p in paths)
    log = changelog_markdown(tr)
    assert "| skills | reordered |" in log
    assert f"{tr.score_before} → {tr.score_after}" in log


def test_write_cover_letter(tmp_path: Path, resume_ru, vacancy_ru) -> None:
    cl = cover_letter(resume_ru, vacancy_ru, MockLLM())
    p = write_cover_letter(cl, tmp_path / "letter.md")
    assert "Здравствуйте" in p.read_text(encoding="utf-8")
