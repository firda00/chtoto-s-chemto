"""Экспорт адаптированного резюме и письма в Markdown и DOCX (ATS-friendly: одна колонка)."""

from __future__ import annotations

import re
from pathlib import Path

from .models import CoverLetter, TailoredResume
from .tailor import SECTION_TITLES

_BULLET = re.compile(r"^\s*[-*•]\s+")
_SUBHEAD = re.compile(r"^\s*#{3,6}\s+")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _titles(tr: TailoredResume) -> dict[str, str]:
    lang = "ru" if tr.platform == "hh" else tr.language
    titles = dict(SECTION_TITLES[lang])
    if tr.platform == "linkedin":
        titles["summary"] = "About" if lang == "en" else "Обо мне"
    return titles


def _strip_md_heading(line: str) -> str:
    return re.sub(r"^\s*#{1,6}\s*", "", line).strip()


def to_markdown(tr: TailoredResume) -> str:
    titles = _titles(tr)
    parts: list[str] = []
    for key, body in tr.sections.items():
        if key == "contacts":
            lines = body.splitlines()
            if lines:
                first = _strip_md_heading(lines[0])
                rest = "\n".join(lines[1:]).strip()
                block = f"# {first}" + (f"\n\n{rest}" if rest else "")
                if tr.platform == "linkedin":
                    parts.append(f"## {titles['contacts']}\n\n{body.strip()}")
                else:
                    parts.insert(0, block)
            continue
        parts.append(f"## {titles.get(key, key.title())}\n\n{body.strip()}")
    return "\n\n".join(parts).strip() + "\n"


def changelog_markdown(tr: TailoredResume) -> str:
    lines = [
        f"# Changelog: адаптация под «{tr.vacancy_title}» ({tr.platform})",
        "",
        f"ATS Score: {tr.score_before} → {tr.score_after}",
        "",
        "| Раздел | Действие | Описание |",
        "|---|---|---|",
    ]
    for ch in tr.changelog:
        lines.append(f"| {ch.section} | {ch.action} | {ch.detail.replace('|', '/')} |")
    if not tr.changelog:
        lines.append("| — | unchanged | Изменений нет |")
    if tr.warnings:
        lines += ["", "## Предупреждения", ""] + [f"- {w}" for w in tr.warnings]
    return "\n".join(lines) + "\n"


def write_markdown(tr: TailoredResume, path: str | Path) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(to_markdown(tr), encoding="utf-8")
    return p


def _add_runs(paragraph, text: str) -> None:
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_body(doc, body: str) -> None:
    for raw in body.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if _BULLET.match(line):
            _add_runs(doc.add_paragraph(style="List Bullet"), _BULLET.sub("", line))
        elif _SUBHEAD.match(line):
            p = doc.add_paragraph()
            p.add_run(_SUBHEAD.sub("", line).replace("**", "")).bold = True
        else:
            _add_runs(doc.add_paragraph(), line.strip())


def write_docx(tr: TailoredResume, path: str | Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    titles = _titles(tr)
    contacts = tr.sections.get("contacts", "")
    if contacts:
        lines = contacts.splitlines()
        doc.add_heading(_strip_md_heading(lines[0]), level=0)
        _add_body(doc, "\n".join(lines[1:]))
    for key, body in tr.sections.items():
        if key == "contacts":
            continue
        doc.add_heading(titles.get(key, key.title()), level=1)
        _add_body(doc, body)
    doc.save(str(p))
    return p


def write_cover_letter(cl: CoverLetter, path: str | Path) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(cl.text.strip() + "\n", encoding="utf-8")
    return p


def write_all(tr: TailoredResume, out_dir: str | Path, stem: str = "resume") -> list[Path]:
    out = Path(out_dir)
    stem = f"{stem}_{tr.platform}"
    return [
        write_markdown(tr, out / f"{stem}.md"),
        write_docx(tr, out / f"{stem}.docx"),
        _write_text(out / f"{stem}_changelog.md", changelog_markdown(tr)),
    ]


def _write_text(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path
